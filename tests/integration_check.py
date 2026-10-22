"""Opt-in live checks: python tests/integration_check.py [--live-chat]."""
import argparse
import asyncio
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
import config
from docx import Document
from PIL import Image
import pymupdf
from services.comparison import compare_files
from rag.rag_pipeline import GraphRAGPipeline


def fixtures(folder):
    folder.mkdir(parents=True, exist_ok=True)
    image = folder / "chart.png"
    Image.new("RGB", (100, 70), "navy").save(image)
    paths = []
    for index in (1, 2):
        budget, month = ("12", "June") if index == 1 else ("18", "September")
        text = f"The Harbor project has a budget of {budget} million dollars. Delivery is scheduled for {month}."
        document = Document()
        document.add_heading("Harbor project review", 0)
        document.add_paragraph(text)
        table = document.add_table(rows=2, cols=2)
        for cell, value in zip((table.cell(0,0),table.cell(0,1),table.cell(1,0),table.cell(1,1)),
                               ("Measure", "Value", "Budget", budget + " million")):
            cell.text = value
        document.add_picture(str(image))
        document.core_properties.title = "Harbor project review"
        document.core_properties.author = "Operations"
        word = folder / f"review-{index}.docx"
        document.save(word)
        pdf_path = folder / f"review-{index}.pdf"
        with pymupdf.open() as pdf:
            page = pdf.new_page()
            page.insert_text((72,72), "Harbor project review", fontsize=18)
            page.insert_textbox(pymupdf.Rect(72,100,520,160), text, fontsize=12)
            page.insert_image(pymupdf.Rect(72,200,172,270), filename=str(image))
            pdf.set_metadata({"title":"Harbor project review","author":"Operations"})
            pdf.save(pdf_path)
        paths.append((word, pdf_path))
    return paths


async def main(live_chat=False):
    folder = ROOT / ".cache" / "validation"
    original, revised = fixtures(folder)
    result = await compare_files(str(original[1]), str(revised[0]), extended=True)
    assert not result.warnings, result.warnings
    assert result.report.image_score.score > 0.99
    assert result.report.text_score.score > 0.4
    assert abs(sum(result.report.weights_used.values()) - 1) < 1e-6
    (folder / "comparison.json").write_text(result.report.model_dump_json(indent=2), encoding="utf-8")
    print("PDF_DOCX_COMPARISON_OK", result.report.overall_score, flush=True)
    identical = await compare_files(str(original[0]), str(original[0]), extended=True)
    assert identical.report.overall_score > 0.99, identical.report.overall_score
    print("IDENTICAL_DOCX_COMPARISON_OK", flush=True)
    pipeline = GraphRAGPipeline(groq_api_key=config.GROQ_API_KEY, rerank=True)
    state = pipeline.ingest(*result.raw_documents)
    try:
        sources, sufficient = pipeline.retrieve("What changed in the Harbor project budget and delivery schedule?", state)
        assert sufficient
        assert {source["doc_id"] for source in sources} == {"doc1", "doc2"}
        assert any(source["page"] == 1 for source in sources if source["doc_id"] == "doc1")
        print("HYBRID_RERANK_GRAPH_RETRIEVAL_OK", len(sources), flush=True)
        if live_chat:
            assert config.GROQ_API_KEY, "Configure GROQ_API_KEY before live chat checks"
            answer = "".join(pipeline.query("Compare the budget and delivery schedule in both documents.", state, stream=True))
            assert "12" in answer and "18" in answer, answer
            assert "June" in answer and "September" in answer, answer
            cited = set(re.findall(r"D[12]-C\d+", answer))
            allowed = {source["source_id"] for source in pipeline.last_sources}
            assert cited and cited.issubset(allowed), (cited, allowed)
            (folder / "chat-answer.md").write_text(answer, encoding="utf-8")
            print("LIVE_STREAMED_CITED_CHAT_OK", flush=True)
        (folder / "retrieval.json").write_text(json.dumps({"sources":sources,"trace":pipeline.last_trace},indent=2), encoding="utf-8")
    finally:
        pipeline.close(state)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--live-chat", action="store_true")
    asyncio.run(main(parser.parse_args().live_chat))
