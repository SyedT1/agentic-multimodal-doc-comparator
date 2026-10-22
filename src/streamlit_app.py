"""Document comparison workspace. Heavy analysis modules load only when needed."""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent))
import config
import asyncio
import hashlib
import json
from html import escape
import streamlit as st
from utils.file_handler import save_uploaded_file
from utils.diff import diff_html, diff_summary, document_text, changed_values

st.set_page_config(page_title="Folio · Document intelligence", page_icon="◫", layout="wide",
                   initial_sidebar_state="auto")
STYLE = """
<style>
:root {font-family:'DM Sans',sans-serif;color:#172b36}
.stApp{background:#f5f7fa}
[data-testid="stSidebar"]{background:#fff;border-right:1px solid #e4e9ee}
[data-testid="stSidebar"] h1{font-size:1.6rem}
.block-container{max-width:1450px;padding-top:2rem;padding-bottom:3rem}
h1,h2,h3{font-family:'Manrope','DM Sans',sans-serif!important;letter-spacing:-.035em}
h1{font-size:2.15rem!important;font-weight:800!important}
h2{font-size:1.4rem!important}h3{font-size:1.1rem!important}
[data-testid="stVerticalBlockBorderWrapper"]>div{border-color:#e1e7ec!important;border-radius:16px!important}
[data-testid="stMetric"]{background:#fff;border:1px solid #e3e9ee;border-radius:12px;padding:16px}
[data-testid="stMetricValue"]{font-family:'Manrope',sans-serif;font-size:1.85rem}
.stButton>button[kind="primary"]{background:#087f78;color:white;border:0;border-radius:10px;font-weight:600;min-height:44px}
.stButton>button[kind="secondary"]{border-radius:9px;border-color:#dce4e9}
[data-testid="stFileUploader"]{background:#fafcfd;border-radius:12px}
[data-baseweb="tab-list"]{gap:1.6rem;border-bottom:1px solid #dde5eb}
[data-baseweb="tab"]{font-weight:600;padding:12px 0}
.brand{display:flex;align-items:center;gap:10px;margin:5px 0 30px;font:800 27px 'Manrope',sans-serif}
.brand-icon{background:#087f78;color:white;border-radius:10px;padding:2px 10px}
.eyebrow{font-size:11px;font-weight:700;letter-spacing:.15em;text-transform:uppercase;color:#69818c;margin-bottom:9px}
.chip{display:inline-block;padding:6px 10px;border-radius:30px;font-size:12px;font-weight:600;background:#e0f2ef;color:#086b66}
.hero{background:#102f3a;color:#fff;border-radius:18px;padding:30px 32px;margin:12px 0 24px}
.hero h2{color:white;margin:0 0 8px!important;font-size:1.5rem!important}
.hero p{color:#b9d1d8;max-width:650px;margin:0;font-size:14px;line-height:1.7}
.step{font-size:11px;letter-spacing:.1em;color:#087f78;font-weight:700}
.score{font:800 58px 'Manrope',sans-serif;color:#087f78;letter-spacing:-.05em}
.small-muted{color:#6d8190;font-size:12px}
.preview{white-space:pre-wrap;font:13px/1.75 'DM Sans',sans-serif}
.footer{border-top:1px solid #e1e7ec;margin-top:28px;padding-top:16px;color:#738792;font-size:12px}
@media(max-width:700px){.block-container{padding:1rem}.hero{padding:22px}h1{font-size:1.65rem!important}}
</style>
"""

def _init():
    for name, value in {"result":None,"chat_pipeline":None,"chat_state":None,"messages":[],
                        "batch_results":[],"last_settings":None,"upload_signature":None}.items():
        if name not in st.session_state:
            st.session_state[name] = value

def _clear_chat():
    pipeline, state = st.session_state.chat_pipeline, st.session_state.chat_state
    if pipeline and state:
        pipeline.close(state)
    st.session_state.chat_pipeline = st.session_state.chat_state = None
    st.session_state.messages = []

def sidebar():
    with st.sidebar:
        st.markdown('<div class="brand"><span class="brand-icon">◫</span> folio<span style="color:#91a3ae;font-size:11px">WORKSPACE</span></div>', unsafe_allow_html=True)
        navigation = st.radio("Workspace", ["Compare documents", "Batch review", "Architecture"], label_visibility="collapsed")
        st.divider()
        st.markdown("##### Analysis settings")
        extended = st.toggle("Visual & structural analysis", value=True,
            help="Include images, document structure, and metadata alongside text and tables.")
        profile = st.selectbox("Scoring profile", ["Balanced", "Content first", "Custom"])
        weights = dict(config.MODALITY_WEIGHTS if extended else config.MODALITY_WEIGHTS_PHASE1)
        if profile == "Content first":
            weights = {"text":0.7,"table":0.15,"image":0.05,"layout":0.05,"metadata":0.05} if extended else {"text":0.85,"table":0.15}
        if profile == "Custom":
            with st.expander("Adjust comparison weights", expanded=True):
                weights = {name: st.slider(name.title(), 0.0, 1.0, float(value), 0.05, key="weight_"+name) for name,value in weights.items()}
        total = sum(weights.values())
        if total <= 0:
            st.warning("Choose at least one nonzero weight. Balanced weights are applied until then.")
            weights = dict(config.MODALITY_WEIGHTS if extended else config.MODALITY_WEIGHTS_PHASE1)
        else:
            weights = {key:value/total for key,value in weights.items()}
        with st.expander("Answer quality"):
            rerank = st.toggle("Rerank evidence", value=True, help="A local model checks question–passage relevance. Downloads once.")
            top_k = st.slider("Evidence passages", 4, 12, 6, 2)
            st.caption("Search combines meaning, keywords, and related passages. Weak evidence triggers one broader search.")
        st.divider()
        st.markdown('<span class="chip">● Local document analysis</span>', unsafe_allow_html=True)
        st.caption("Chat connected" if config.GROQ_API_KEY else "Add a Groq key to enable chat")
        st.caption("PDF & DOCX · Maximum 50 MB per file")
    return navigation, weights, extended, rerank, top_k

def _run(path1, path2, weights, extended):
    from services.comparison import compare_files
    bar = st.progress(0, text="Preparing comparison")
    try:
        result = asyncio.run(compare_files(path1, path2, weights, extended,
            lambda fraction,message: bar.progress(fraction, text=message)))
        _clear_chat()
        st.session_state.result = result
        st.session_state.last_settings = (weights,extended)
        return result
    finally:
        bar.empty()

def _demos():
    from docx import Document
    from PIL import Image
    folder = config.CACHE_DIR / "examples"
    folder.mkdir(parents=True, exist_ok=True)
    image = folder / "mark.png"
    Image.new("RGB",(80,40),"#087f78").save(image)
    paths = []
    for index in (1,2):
        doc = Document()
        doc.add_heading("Harbor project · Quarterly review",0)
        doc.add_heading("Executive summary",1)
        doc.add_paragraph("The Harbor project expands regional delivery capacity. This review covers budget, milestones, and staffing.")
        doc.add_heading("Budget and delivery",1)
        doc.add_paragraph("The approved budget is $12 million. Delivery is scheduled for June." if index == 1
                         else "The revised budget is $18 million. Delivery is scheduled for September.")
        table = doc.add_table(rows=1,cols=2)
        table.rows[0].cells[0].text="Measure";table.rows[0].cells[1].text="Target"
        for label,value in (("Budget", "$12 million" if index == 1 else "$18 million"),
                            ("Team size","24" if index == 1 else "30")):
            cells=table.add_row().cells;cells[0].text=label;cells[1].text=value
        doc.add_picture(str(image))
        doc.core_properties.title="Harbor quarterly review"
        doc.core_properties.author="Operations team"
        path=folder / ("harbor-original.docx" if index == 1 else "harbor-revised.docx")
        doc.save(path);paths.append(str(path))
    return paths

def uploads(weights, extended):
    st.markdown('<div class="eyebrow">DOCUMENT INTELLIGENCE</div>', unsafe_allow_html=True)
    left,right=st.columns([4,1])
    with left:
        st.title("See what changed.")
        st.caption("A clear view of what your documents share, where they differ, and why it matters.")
    with right:
        demo=st.button("Try sample documents", width="stretch")
    st.markdown('<div class="hero"><div class="eyebrow" style="color:#77c8bd">YOUR COMPARISON WORKSPACE</div><h2>Two documents. One complete picture.</h2><p>Compare the words, tables, visuals, and structure. Explore the changes side by side, then ask questions grounded in the source.</p></div>', unsafe_allow_html=True)
    columns=st.columns(2,gap="large")
    files=[]
    for index,col in enumerate(columns):
        with col, st.container(border=True):
            st.markdown(f'<div class="step">0{index+1} / {"REFERENCE" if index == 0 else "COMPARISON"}</div>', unsafe_allow_html=True)
            st.subheader("Original document" if index == 0 else "Revised document")
            file=st.file_uploader("Choose a PDF or Word document",type=["pdf","docx"],key=f"document_{index}",max_upload_size=50)
            files.append(file)
            st.caption(f"{file.name} · {file.size/1024:,.0f} KB" if file else "Drop a file here to get started.")
    current_signature = tuple(hashlib.sha256(file.getbuffer()).hexdigest() if file else "" for file in files)
    result=st.session_state.result
    changed = result is not None and ((any(current_signature) and current_signature != st.session_state.upload_signature)
                                     or st.session_state.last_settings != (weights,extended))
    if changed:
        st.warning("Documents or settings have changed. Compare again to update the results and chat sources.")
    button_col,hint_col=st.columns([1,2])
    with button_col:
        compare=st.button("Compare documents",type="primary",width="stretch",disabled=not all(files))
    with hint_col:
        st.caption("Scoring uses only the content available in your files. Similarity is not a factual accuracy check.")
    if demo or compare:
        try:
            paths=_demos() if demo else [save_uploaded_file(file) for file in files]
            with st.spinner("Analyzing documents. Models download only on first use."):
                _run(*paths,weights,extended)
            st.session_state.upload_signature=current_signature
            changed=False
            st.success("Comparison ready. Explore the results below.")
        except Exception as exc:
            st.error(f"Could not complete the comparison: {exc}")
    return changed

def overview(result):
    report=result.report
    st.subheader("Comparison overview")
    for warning in result.warnings:
        st.warning(warning)
    left,right=st.columns([1,2],gap="large")
    with left,st.container(border=True):
        st.markdown('<div class="eyebrow">OVERALL SIMILARITY</div>',unsafe_allow_html=True)
        st.markdown(f'<div class="score">{report.overall_score:.0%}</div>',unsafe_allow_html=True)
        st.progress(report.overall_score)
        st.caption("A weighted measure of shared content. See the evidence before drawing conclusions.")
    with right:
        changes=diff_summary(*(document_text(document) for document in result.raw_documents))
        cols=st.columns(3)
        for col,(label,count) in zip(cols,changes.items()):
            col.metric(label.title()+" lines",count)
        st.caption(f"{report.doc1_name}  →  {report.doc2_name}")
        st.caption("Word documents have logical sections; physical page numbers are available for PDFs.")
    scores=[]
    for name,label in (("text","Text"),("table","Tables"),("image","Images"),("layout","Structure"),("metadata","Metadata")):
        score=getattr(report,name+"_score")
        available=score is not None and score.details.get("available",True)
        scores.append({"Content":label,"Similarity":score.score if available else None,
                       "Weight":report.weights_used.get(name,0),"Status":"Compared" if available else "Not available"})
    st.dataframe(scores,hide_index=True,width="stretch",
        column_config={"Similarity":st.column_config.ProgressColumn(format="percent",min_value=0,max_value=1),
                       "Weight":st.column_config.NumberColumn(format="percent")})
    st.caption("Unavailable content is excluded. Content present in only one document counts as a difference.")
    removed,added=changed_values(*(document_text(document) for document in result.raw_documents))
    if removed or added:
        st.warning("Values or dates changed. High semantic similarity can still contain important factual differences.")
        first,second=st.columns(2)
        first.write("**Only in original:** " + (", ".join(removed) or "None"))
        second.write("**Only in revised:** " + (", ".join(added) or "None"))
    st.markdown("#### Matching passages")
    matches=[m for m in report.matched_sections if m["type"]=="text"]
    if not matches:
        st.info("No text passages exceeded the matching threshold.")
    for index,match in enumerate(matches):
        with st.expander(f"Passage {index+1} · {match['similarity']:.0%} semantic similarity",expanded=index==0):
            columns=st.columns(2)
            for column,key,label in zip(columns,("doc1_content","doc2_content"),("Original","Revised")):
                with column:
                    st.caption(label)
                    st.write(match[key])
    st.download_button("Download comparison report",json.dumps(report.model_dump(mode="json"),indent=2),
                       "comparison-report.json","application/json",width="content")

def documents(result):
    raw1,raw2=result.raw_documents
    st.subheader("Read the documents")
    columns=st.columns(2,gap="large")
    for index,(col,raw,processed) in enumerate(zip(columns,result.raw_documents,result.processed_documents)):
        with col,st.container(border=True):
            st.markdown("**"+raw.filename+"**")
            if raw.file_type=="pdf":
                number=st.number_input("Page",1,max(1,raw.total_pages),key=f"preview_page_{index}")
                import pymupdf
                with pymupdf.open(raw.metadata["file_path"]) as pdf:
                    page=pdf[int(number)-1]
                    st.image(page.get_pixmap(matrix=pymupdf.Matrix(1.2,1.2)).tobytes("png"),width="stretch")
                text=next((p["text"] for p in raw.pages if p["page_num"]==number),"")
            else:
                st.caption("Word content · logical sections")
                text=raw.raw_text
            with st.container(height=240):
                st.text(text[:30000])
            if processed.tables:
                with st.expander(f"Tables ({len(processed.tables)})"):
                    for table in raw.raw_tables[:10]:
                        st.dataframe(table["data"],hide_index=True,width="stretch")
            if processed.images:
                with st.expander(f"Images ({len(processed.images)})"):
                    st.image([image.image for image in processed.images[:8]],width=140)
    st.subheader("Text changes")
    st.caption("Green: added · Red: removed · Amber: edited. The viewer covers the first 5,000 lines of extracted text.")
    markup=diff_html(document_text(raw1),document_text(raw2),raw1.filename,raw2.filename)
    st.iframe(markup,height=430)
    st.download_button("Download text diff",markup,"document-changes.html","text/html")

def _sources(sources):
    for source in sources:
        location=f"page {source['page']}" if source.get("page") else "Word section"
        with st.expander(f"[{source['source_id']}] {source.get('filename','Document')} · {location}"):
            st.write(source["text"])

def chat(result,stale,rerank,top_k):
    st.subheader("Ask your documents")
    st.caption("Answers include source references you can inspect. Ask about a change, a number, or a specific section.")
    if stale:
        st.info("Compare the updated documents before asking questions.")
        return
    pipeline=st.session_state.chat_pipeline
    signature=(rerank,top_k)
    if pipeline and getattr(pipeline,"ui_settings",None)!=signature:
        _clear_chat();pipeline=None
    if not pipeline:
        if st.button("Prepare document chat",type="primary"):
            try:
                from rag.rag_pipeline import GraphRAGPipeline
                with st.spinner("Preparing searchable evidence"):
                    pipeline=GraphRAGPipeline(groq_api_key=config.GROQ_API_KEY,rerank=rerank,top_k_vector=top_k)
                    state=pipeline.ingest(*result.raw_documents)
                    pipeline.ui_settings=signature
                    st.session_state.chat_pipeline=pipeline
                    st.session_state.chat_state=state
            except Exception as exc:
                st.error(f"Could not prepare chat: {exc}")
                return
        else:
            st.info("Prepare the documents once to start asking questions. Search works locally; answers use Groq.")
            return
    pipeline=st.session_state.chat_pipeline
    state=st.session_state.chat_state
    c1,c2=st.columns([4,1])
    with c1:
        st.caption(f"{state.stats['total_vectors']} searchable passages across both documents")
    with c2:
        if st.button("Clear conversation"):
            pipeline.reset_chat();st.session_state.messages=[];st.rerun()
    if not config.GROQ_API_KEY:
        st.info("Add GROQ_API_KEY to your local .env file and restart to enable answers.")
        return
    for message in st.session_state.messages:
        with st.chat_message(message["role"]):
            st.markdown(message["content"])
            if message.get("sources"):
                _sources(message["sources"])
    question=st.chat_input("What changed in the budget or timeline?")
    if question:
        st.session_state.messages.append({"role":"user","content":question})
        with st.chat_message("user"):
            st.write(question)
        with st.chat_message("assistant"):
            try:
                with st.spinner("Finding and checking evidence"):
                    response=pipeline.query(question,state,stream=True)
                answer=st.write_stream(response)
                sources=[dict(source) for source in pipeline.last_sources]
                st.session_state.messages.append({"role":"assistant","content":answer,"sources":sources})
                _sources(sources)
                with st.expander("How the answer was found"):
                    for step in pipeline.last_trace:
                        st.write(step)
            except Exception as exc:
                st.error(f"The answer could not be completed: {exc}. Please retry.")
                st.session_state.messages.pop()

def batch(weights,extended):
    st.markdown('<div class="eyebrow">COLLECTION REVIEW</div>',unsafe_allow_html=True)
    st.title("Find the closest match.")
    st.caption("Compare one reference against up to 10 candidates. Each result uses the same analysis settings.")
    left,right=st.columns(2)
    with left:
        reference=st.file_uploader("Reference document",type=["pdf","docx"],key="batch_reference",max_upload_size=50)
    with right:
        candidates=st.file_uploader("Candidate documents",type=["pdf","docx"],accept_multiple_files=True,key="batch_candidates",max_upload_size=50)
    if candidates and len(candidates)>10:
        st.warning("Choose at most 10 candidate documents.");return
    if st.button("Review collection",type="primary",disabled=not reference or not candidates):
        from services.comparison import compare_files
        path=save_uploaded_file(reference)
        rows=[]
        progress=st.progress(0)
        for index,candidate in enumerate(candidates):
            try:
                result=asyncio.run(compare_files(path,save_uploaded_file(candidate),weights,extended))
                rows.append({"Document":candidate.name,"Similarity":result.report.overall_score,"Status":"Complete"})
            except Exception as exc:
                rows.append({"Document":candidate.name,"Similarity":None,"Status":str(exc)})
            progress.progress((index+1)/len(candidates))
        st.session_state.batch_results=sorted(rows,key=lambda row:row["Similarity"] if row["Similarity"] is not None else -1,reverse=True)
    if st.session_state.batch_results:
        import pandas as pd
        frame=pd.DataFrame(st.session_state.batch_results)
        st.dataframe(frame,hide_index=True,width="stretch",column_config={"Similarity":st.column_config.ProgressColumn(format="percent",min_value=0,max_value=1)})
        st.download_button("Download batch results",frame.to_csv(index=False),"batch-results.csv","text/csv")

def architecture():
    st.markdown('<div class="eyebrow">UNDER THE HOOD</div>',unsafe_allow_html=True)
    st.title("Evidence before answers.")
    st.write("Comparison runs locally. Chat uses a bounded retrieval workflow and sends only the selected passages to Groq.")
    st.markdown("""
1. **Read** PDF and Word text, tables, images, structure, and metadata.
2. **Compare** both directions and normalize weights over available content.
3. **Index** page-aware passages in a separate Chroma collection for each comparison.
4. **Plan** a standalone question and which documents to search.
5. **Retrieve** semantic and BM25 keyword matches, then expand through related document passages.
6. **Check** relevance with a local reranker. Broaden the search at most once when evidence is weak.
7. **Answer** with stable source references, or explain that evidence is insufficient.
""")
    st.info("The document graph links passages by sequence, shared terms, and similarity. It is not a database of verified facts.")
    st.markdown("#### Tooling review")
    audit = Path(__file__).resolve().parents[1] / "docs" / "TOOLING_AUDIT_2026.md"
    if audit.exists():
        with st.expander("Read the October 2026 tooling audit"):
            st.markdown(audit.read_text(encoding="utf-8"))
    st.caption("Scanned PDFs currently need an OCR text layer. Physical pagination is not inferred for Word files.")

async def process_and_compare(file1_path,file2_path,weights,enable_phase2=False):
    from services.comparison import compare_files
    result=await compare_files(file1_path,file2_path,weights,enable_phase2)
    return result.report,*result.raw_documents

def main():
    st.markdown(STYLE,unsafe_allow_html=True)
    _init()
    navigation,weights,extended,rerank,top_k=sidebar()
    if navigation=="Batch review":
        batch(weights,extended)
    elif navigation=="Architecture":
        architecture()
    else:
        stale=uploads(weights,extended)
        result=st.session_state.result
        if result:
            overview_tab,documents_tab,chat_tab=st.tabs(["Overview","Documents & changes","Ask documents"],
                                                       key="comparison_view",on_change="rerun")
            if overview_tab.open:
                with overview_tab:overview(result)
            if documents_tab.open:
                with documents_tab:documents(result)
            if chat_tab.open:
                with chat_tab:chat(result,stale,rerank,top_k)
        else:
            st.markdown("#### Your review starts here")
            st.caption("Upload two documents or try the sample pair. Results, document previews, changes, and chat will appear in this workspace.")
    st.markdown('<div class="footer">◫ folio &nbsp; / &nbsp; Document intelligence workspace <span style="float:right">Local analysis · Source-grounded answers</span></div>',unsafe_allow_html=True)

if __name__=="__main__":
    main()
